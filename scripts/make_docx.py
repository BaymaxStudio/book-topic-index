#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""命中明细 → 「勘校索引」版式 Word。

设计意图（这不是普通报告模板）：
  文档的职责是让人快速定位并引用，所以正文区做成左栏页码导轨 + 右栏引文，
  检索词以朱色（朱批）标出——它本身就是一份在书上做批注的索引。

用法:
  python3 make_docx.py --hits <命中明细.json> --out <索引.docx> --config book.json
"""
import argparse, json, os, re
from collections import Counter
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INK      = RGBColor(0x16, 0x18, 0x1D)   # 墨
INDIGO   = RGBColor(0x1F, 0x3A, 0x5F)   # 靛：结构色
VERM     = RGBColor(0xA8, 0x32, 0x2A)   # 朱：只给关键词
SLATE    = RGBColor(0x6B, 0x72, 0x80)
SILVER_C = RGBColor(0xD9, 0xDC, 0xE1)
SILVER   = "D9DCE1"
HEI, SONG, LAT = "黑体", "宋体", "Times New Roman"

# OOXML 子元素次序是固定的，手工插入时必须按序，否则严格校验器会报 schema 错误
PPR_AFTER_PBDR = ('w:shd','w:tabs','w:spacing','w:ind','w:jc','w:rPr','w:sectPr','w:pPrChange')
RPR_AFTER_SPACING = ('w:w','w:kern','w:position','w:sz','w:szCs','w:u','w:effect','w:bdr',
                     'w:shd','w:fitText','w:vertAlign','w:rtl','w:cs','w:lang','w:eastAsianLayout')
TCPR_AFTER_BORDERS = ('w:shd','w:noWrap','w:tcMar','w:textDirection','w:tcFitText',
                      'w:vAlign','w:hideMark','w:tcPrChange')
TBLPR_AFTER_LAYOUT = ('w:tblCellMar','w:tblLook','w:tblCaption','w:tblDescription','w:tblPrChange')
BORDER_ORDER = ('top','left','bottom','right','insideH','insideV','tl2br','tr2bl')

def _rfonts(run, ea, latin=LAT):
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rPr.insert(0, rf)
    rf.set(qn('w:ascii'), latin); rf.set(qn('w:hAnsi'), latin); rf.set(qn('w:eastAsia'), ea)

def R(run, ea=SONG, size=10.5, color=INK, bold=False, caps=0):
    run.font.size = Pt(size); run.font.color.rgb = color; run.bold = bold
    _rfonts(run, ea)
    if caps:
        rPr = run._element.get_or_add_rPr()
        sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), str(int(caps)))
        rPr.insert_element_before(sp, *RPR_AFTER_SPACING)
    return run

def P(container, before=0, after=0, line=1.5, align=None):
    p = container.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    if line: pf.line_spacing = line
    if align is not None: p.alignment = align
    return p

def para_border(paragraph, side='bottom', color=SILVER, sz=6, space=0):
    """在正文段落上画线，避免用空段落当分隔线"""
    pPr = paragraph._p.get_or_add_pPr()
    b = pPr.find(qn('w:pBdr'))
    if b is None:
        b = OxmlElement('w:pBdr'); pPr.insert_element_before(b, *PPR_AFTER_PBDR)
    el = OxmlElement('w:' + side)
    el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(sz))
    el.set(qn('w:space'), str(space)); el.set(qn('w:color'), color)
    b.append(el); return paragraph

def set_cell_border(cell, **kw):
    """一个单元格只能有一个 tcBorders，多边必须合并写入"""
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        b = OxmlElement('w:tcBorders'); tcPr.insert_element_before(b, *TCPR_AFTER_BORDERS)
    for side in BORDER_ORDER:
        if side not in kw: continue
        spec = kw[side]; el = OxmlElement('w:' + side)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(spec.get('sz', 6)))
        el.set(qn('w:space'), str(spec.get('space', 0)))
        el.set(qn('w:color'), spec.get('color', SILVER)); b.append(el)

def bookmark(paragraph, name, bid):
    p = paragraph._p
    st = OxmlElement('w:bookmarkStart'); st.set(qn('w:id'), str(bid)); st.set(qn('w:name'), name)
    en = OxmlElement('w:bookmarkEnd'); en.set(qn('w:id'), str(bid))
    pPr = p.find(qn('w:pPr'))
    if pPr is not None: pPr.addnext(st)     # 必须落在 pPr 之后
    else: p.insert(0, st)
    p.append(en)

def add_link(paragraph, text, anchor, size=9, color=INDIGO, ea=HEI):
    hl = OxmlElement('w:hyperlink'); hl.set(qn('w:anchor'), anchor)
    r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts')
    rf.set(qn('w:ascii'), LAT); rf.set(qn('w:hAnsi'), LAT); rf.set(qn('w:eastAsia'), ea)
    rPr.append(rf); rPr.append(OxmlElement('w:b'))
    c = OxmlElement('w:color'); c.set(qn('w:val'), str(color)); rPr.append(c)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size * 2))); rPr.append(sz)
    r.append(rPr)
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
    r.append(t); hl.append(r); paragraph._p.append(hl)

def page_field(p, size=8.5, color=SLATE):
    r = p.add_run(); R(r, HEI, size, color)
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = 'PAGE'
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end')
    r._r.append(f1); r._r.append(it); r._r.append(f2)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--hits", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--config")
    ap.add_argument("--prepare", help="pdf_lines.py 产出的 prepare.json，用于判断是否 OCR")
    a = ap.parse_args()
    cfg = {"title": "未命名", "edition": "", "author": "", "term_label": ""}
    if a.config and os.path.exists(a.config):
        cfg.update(json.load(open(a.config, encoding="utf-8")))
    rows = json.load(open(a.hits, encoding="utf-8"))
    KW = re.compile("|".join(cfg["terms"])) if cfg.get("terms") else re.compile(r"(?!)")

    body = [r for r in rows if r["类型"] != "目录"]
    toc  = [r for r in rows if r["类型"] == "目录"]
    # 分组键三级回退：书内页码 → PDF 页 → 无页码。
    # 只认书内页码会让"仅 Word"输入的条目整批丢失（无 PDF 就没有印刷页码）。
    def gkey(r):
        if r.get("书内页码") is not None: return ("页", r["书内页码"])
        if r.get("PDF页") is not None:     return ("PDF", r["PDF页"])
        return ("无", 0)

    groups = {}
    for r in body:
        groups.setdefault(gkey(r), []).append(r)
    keys = sorted(groups, key=lambda k: ({"页": 0, "PDF": 1, "无": 2}[k[0]], k[1]))
    pages = [k[1] for k in keys if k[0] == "页"]
    has_pages = bool(pages)
    total_occ = sum(r["命中次数"] for r in rows)
    offs = Counter(r["PDF页"] - r["书内页码"] for r in body if r["书内页码"])
    offset = offs.most_common(1)[0][0] if offs else 0

    title = cfg["title"] + ("（%s）" % cfg["edition"] if cfg["edition"] else "")
    doc = Document()
    st = doc.styles["Normal"]; st.font.size = Pt(10.5); st.font.name = LAT
    st.element.rPr.rFonts.set(qn('w:eastAsia'), SONG)
    meta = doc.styles.add_style('IndexMeta', WD_STYLE_TYPE.PARAGRAPH)
    meta.base_style = doc.styles['Normal']
    meta.font.size = Pt(10.5); meta.font.name = LAT
    meta.element.rPr.rFonts.set(qn('w:eastAsia'), SONG)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.3)
    sec.top_margin = Cm(2.4); sec.bottom_margin = Cm(2.2)

    p = P(doc, before=36, after=14, line=1.0)
    R(p.add_run("文献检索索引"), HEI, 9, INDIGO, True, caps=90)
    R(p.add_run("　／　LITERATURE INDEX"), HEI, 8, SLATE, False, caps=40)

    p = P(doc, after=2, line=1.15)
    R(p.add_run(title), HEI, 21, INK, True)

    p = P(doc, after=16, line=1.2)
    who = "著者　%s" % cfg["author"] if cfg["author"] else ""
    R(p.add_run(who), SONG, 10.5, SLATE)
    para_border(p, 'bottom', '1F3A5F', 10, space=8)

    p = P(doc, before=16, after=4, line=1.3)
    R(p.add_run("检索词"), HEI, 9, SLATE, True, caps=50)
    R(p.add_run("　　"), SONG, 10)
    R(p.add_run(cfg["term_label"] or "关键词"), HEI, 15, VERM, True)

    p = P(doc, after=10, line=1.5)
    stat = [(str(len(body)), "命中段落")]
    stat.append((str(len(pages)), "分布页") if has_pages else (str(len(keys)), "命中分组"))
    stat.append((str(total_occ), "全书出现"))
    stat.append((str(offset), "页码偏移") if has_pages else ("—", "页码偏移"))
    for i, (num, lab) in enumerate(stat):
        if i: R(p.add_run("　·　"), HEI, 10, SILVER_C)
        R(p.add_run(num), HEI, 12, INDIGO, True)
        R(p.add_run(" " + lab), SONG, 9.5, SLATE)

    p = P(doc, after=10, line=1.5)
    R(p.add_run("说明　"), HEI, 9, INDIGO, True, caps=30)
    R(p.add_run("页码为书内印刷页码，PDF 为电子文件物理页；" if has_pages
                else "本文档没有可依据的印刷页码，条目不标页码、不编造；"), SONG, 9, SLATE)
    R(p.add_run("朱色字"), SONG, 9, VERM, True)
    mode = ""
    if a.prepare and os.path.exists(a.prepare):
        try: mode = json.load(open(a.prepare, encoding="utf-8")).get("mode", "")
        except Exception: mode = ""
    tail = ("为命中关键词；扫描版引文经 OCR 识别，正式引用请对照原书核校。"
            if mode != "text" else
            "为命中关键词；引文取自 PDF 文字层，未经 OCR。")
    R(p.add_run(tail), SONG, 9, SLATE)

    p = P(doc, after=18, line=1.6)
    if has_pages:
        R(p.add_run("命中页　"), HEI, 8.5, SLATE, True, caps=40)
        for i, pg in enumerate(pages):
            if i: R(p.add_run("　·　"), HEI, 9, SILVER_C)
            add_link(p, str(pg), "pg%s" % pg, size=9)
    else:
        R(p.add_run("排列　"), HEI, 8.5, SLATE, True, caps=40)
        R(p.add_run("按原文出现顺序，未标注页码"), SONG, 9.5, SLATE)
    para_border(p, 'bottom', '1F3A5F', 6, space=10)

    for p_ in doc.paragraphs:
        p_.style = meta

    for key in keys:
        rs = groups[key]
        kind, val = key
        if kind == "页":
            rail_main, rail_sub, anchor = str(val), "PDF %s" % rs[0]["PDF页"], "pg%s" % val
        elif kind == "PDF":
            rail_main, rail_sub, anchor = str(val), "PDF 物理页", "pgpdf%s" % val
        else:
            rail_main, rail_sub, anchor = "—", "页码未提供", "pgnone"
        tbl = doc.add_table(rows=1, cols=2)
        tbl.autofit = False
        tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
        tblPr = tbl._tbl.tblPr
        mar = OxmlElement('w:tblCellMar')
        for side, w in (('top', '150'), ('left', '0'), ('bottom', '190'), ('right', '130')):
            e = OxmlElement('w:' + side); e.set(qn('w:w'), w); e.set(qn('w:type'), 'dxa'); mar.append(e)
        tblPr.insert_element_before(mar, 'w:tblLook','w:tblCaption','w:tblDescription','w:tblPrChange')

        left, right = tbl.rows[0].cells
        left.width = Cm(2.5); right.width = Cm(13.9)
        tbl.columns[0].width = Cm(2.5); tbl.columns[1].width = Cm(13.9)
        left.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        right.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        set_cell_border(left, bottom=dict(color=SILVER, sz=6))
        set_cell_border(right, bottom=dict(color=SILVER, sz=6),
                        left=dict(color=SILVER, sz=4, space=10))

        lp = left.paragraphs[0]
        lp.paragraph_format.space_after = Pt(0); lp.paragraph_format.line_spacing = 1.0
        R(lp.add_run(rail_main), HEI, 17, INDIGO, True)
        bookmark(lp, anchor, abs(hash(anchor)) % 100000)
        lp2 = left.add_paragraph()
        lp2.paragraph_format.space_before = Pt(1); lp2.paragraph_format.line_spacing = 1.0
        R(lp2.add_run(rail_sub), HEI, 7.5, SLATE, caps=20)

        rp = right.paragraphs[0]
        ctx = " ".join(x for x in [rs[0].get("章", ""), rs[0].get("节", "")] if x)
        if ctx:
            rp.paragraph_format.space_after = Pt(6); rp.paragraph_format.line_spacing = 1.15
            rp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            R(rp.add_run(ctx), HEI, 8.5, INDIGO, True, caps=20)

        for i, r in enumerate(rs):
            tp = rp if (i == 0 and not ctx) else right.add_paragraph()
            tp.paragraph_format.space_after = Pt(8)
            tp.paragraph_format.line_spacing = 1.55
            tp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            ind = tp._p.get_or_add_pPr().get_or_add_ind()
            ind.set(qn('w:firstLineChars'), '200')
            if r.get("层级") == "相关":
                R(tp.add_run("〔相关词〕"), HEI, 9, SLATE, True)
            rich(tp, r["段落"], KW)

    if toc:
        p = P(doc, before=22, after=8, line=1.2)
        p.style = meta
        R(p.add_run("目录中的相关条目"), HEI, 12, INK, True)
        para_border(p, 'top', '1F3A5F', 6, space=12)
        for r in toc:
            clean = re.sub(r"[.·•…\s]{2,}\d*\s*$", "", r["段落"]).strip()
            p = P(doc, after=4, line=1.4)
            ind = p._p.get_or_add_pPr().get_or_add_ind()
            ind.set(qn('w:firstLineChars'), '200')
            R(p.add_run("指向第 %s 页　" % (r["指向页码"] or "—")), HEI, 9, INDIGO, True)
            rich(p, clean, KW, size=10)

    ft = sec.footer.paragraphs[0]
    ft.alignment = WD_ALIGN_PARAGRAPH.CENTER
    R(ft.add_run("%s索引　·　第 " % (cfg["term_label"] or "")), HEI, 8.5, SLATE)
    page_field(ft)
    R(ft.add_run(" 页"), HEI, 8.5, SLATE)

    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    doc.save(a.out)
    print("saved: %s | groups=%d entries=%d" % (a.out, len(keys), len(body)))

def rich(p, text, KW, size=10.5, ea=SONG, color=INK):
    pos = 0
    for m in KW.finditer(text):
        if m.start() > pos:
            R(p.add_run(text[pos:m.start()]), ea, size, color)
        R(p.add_run(m.group(0).replace(" ", "").replace("\u3000", "")), ea, size, VERM, True)
        pos = m.end()
    if pos < len(text):
        R(p.add_run(text[pos:]), ea, size, color)
    return p

if __name__ == "__main__":
    main()
